import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

INPUT_PATH = "results/golden_set_reviewed.csv"
OUTPUT_PATH = "results/golden_set_human.xlsx"

INTENTS = [
    "Delivery Issue",
    "Order Status",
    "Refund / Return",
    "Payment / Charges",
    "Account / Login",
    "Prime Membership",
    "Digital Content",
    "Product / Seller",
    "Technical Issue",
    "Customer Support / Escalation",
    "Other"
]

print("Loading Golden Set...")

df = pd.read_csv(INPUT_PATH)

# Take 150 examples
df = df.dropna(subset=["customer_text", "amazon_reply"])
df = df.sample(n=150, random_state=42).reset_index(drop=True)

# Keep useful columns
review_df = pd.DataFrame({
    "Example": range(1, 151),
    "Customer Message": df["customer_text"],
    "Historical Amazon Reply": df["amazon_reply"],
    "Suggested Intent": df.get("intent", ""),
    "Human Intent": "",
    "Notes": ""
})

# Save Excel
review_df.to_excel(OUTPUT_PATH, index=False)

# Add dropdown for Human Intent
wb = load_workbook(OUTPUT_PATH)
ws = wb.active
ws.title = "Golden Set Review"

dropdown = DataValidation(
    type="list",
    formula1='"' + ",".join(INTENTS) + '"',
    allow_blank=True
)

ws.add_data_validation(dropdown)

# Human Intent column = E
dropdown.add("E2:E151")

# Formatting
ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:F151"

ws.column_dimensions["A"].width = 12
ws.column_dimensions["B"].width = 60
ws.column_dimensions["C"].width = 60
ws.column_dimensions["D"].width = 25
ws.column_dimensions["E"].width = 30
ws.column_dimensions["F"].width = 45

for row in ws.iter_rows():
    for cell in row:
        cell.alignment = cell.alignment.copy(
            wrap_text=True,
            vertical="top"
        )

wb.save(OUTPUT_PATH)

print()
print("========================================")
print("GOLDEN SET REVIEW SHEET CREATED")
print("========================================")
print(f"150 examples saved to:")
print(OUTPUT_PATH)
print()
print("Open the Excel file and fill the 'Human Intent' column.")
print("Select the correct intent from the dropdown.")
print("Add notes only when useful.")
print()